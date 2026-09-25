#!/usr/bin/env python3
"""
STEP 6 - Build the final Excel workbook (and optional CSVs) from decisions.

Columns (locked with user):
  ICDO32 Code | ICDO32 Term | ICDO32 Preferred Term | Site-Specific? |
  Site Subset (C-codes) | Site Subset (Labels) | Subset Site Count |
  Code's SMVL-Valid Site Count | Code's SMVL-Valid Sites (organ summary) |
  Code's SMVL-Valid Sites (all C-codes) | Rationale | Source (URL) | Source Reference

Sheets: README | Site-Specific Terms (Yes only) | All Screened Terms (all).
Styling: Calibri; teal header #20808D; layout starts at B2; data header row 5.

Input : decisions.jsonl, ceilings.json, pref_map.json
Output: <out.xlsx>   (CSVs optional via --csv)

Usage:
  python3 06_build_workbook.py decisions.jsonl ceilings.json pref_map.json OUT.xlsx [--csv]
"""
import json, sys, csv
from collections import defaultdict

ORGAN = {
 'C00':'Lip','C01':'Tongue','C02':'Tongue','C03':'Gum','C04':'Mouth','C05':'Palate','C06':'Mouth',
 'C07':'Salivary','C08':'Salivary','C09':'Tonsil','C10':'Oropharynx','C11':'Nasopharynx',
 'C12':'Pyriform sinus','C13':'Hypopharynx','C14':'Pharynx','C15':'Esophagus','C16':'Stomach',
 'C17':'Small intestine','C18':'Colon','C19':'Rectosigmoid','C20':'Rectum','C21':'Anus',
 'C22':'Liver','C23':'Gallbladder','C24':'Biliary','C25':'Pancreas','C26':'GI tract',
 'C30':'Nasal/middle ear','C31':'Sinuses','C32':'Larynx','C33':'Trachea','C34':'Lung',
 'C37':'Thymus','C38':'Heart/mediastinum','C39':'Respiratory','C40':'Bone-limbs','C41':'Bone-other',
 'C42':'Blood/marrow','C44':'Skin','C47':'Peripheral nerves','C48':'Peritoneum/retroperitoneum',
 'C49':'Soft tissue','C50':'Breast','C51':'Vulva','C52':'Vagina','C53':'Cervix','C54':'Uterus',
 'C55':'Uterus','C56':'Ovary','C57':'Female genital','C58':'Placenta','C60':'Penis','C61':'Prostate',
 'C62':'Testis','C63':'Male genital','C64':'Kidney','C65':'Renal pelvis','C66':'Ureter','C67':'Bladder',
 'C68':'Urinary','C69':'Eye','C70':'Meninges','C71':'Brain','C72':'CNS','C73':'Thyroid','C74':'Adrenal',
 'C75':'Endocrine','C76':'Ill-defined','C77':'Lymph nodes','C80':'Unknown primary',
}
def organ(code): return ORGAN.get(code[:3], code[:3])
def labelize(codes): return '; '.join(f'{c}={organ(c)}' for c in codes)
def summarize(codes):
    g = defaultdict(list)
    for c in codes: g[organ(c)].append(c)
    return '; '.join(f'{k} ({len(v)})' for k, v in sorted(g.items()))

COLS = ["ICDO32 Code","ICDO32 Term","ICDO32 Preferred Term","Site-Specific?",
        "Site Subset (C-codes)","Site Subset (Labels)","Subset Site Count",
        "Code's SMVL-Valid Site Count","Code's SMVL-Valid Sites (organ summary)",
        "Code's SMVL-Valid Sites (all C-codes)","Rationale","Source (URL)","Source Reference"]

def rows_from(decisions, ceilings, pref):
    for r in sorted(decisions, key=lambda x: (x["code"], x["term"])):
        code = r["code"]; ceil = sorted(ceilings.get(code, []))
        subset = r.get("subset_codes", [])
        yield {
            "ICDO32 Code": code,
            "ICDO32 Term": r["term"],
            "ICDO32 Preferred Term": pref.get(code, ""),
            "Site-Specific?": r["decision"],
            "Site Subset (C-codes)": ', '.join(subset),
            "Site Subset (Labels)": labelize(subset),
            "Subset Site Count": r.get("subset_count", 0),
            "Code's SMVL-Valid Site Count": len(ceil),
            "Code's SMVL-Valid Sites (organ summary)": summarize(ceil),
            "Code's SMVL-Valid Sites (all C-codes)": ', '.join(ceil),
            "Rationale": r.get("rationale", ""),
            "Source (URL)": r.get("source", ""),
            "Source Reference": r.get("source_name", ""),
        }

def main():
    decisions = [json.loads(l) for l in open(sys.argv[1])]
    ceilings = json.load(open(sys.argv[2]))
    pref = json.load(open(sys.argv[3]))
    out = sys.argv[4]
    want_csv = "--csv" in sys.argv

    all_rows = list(rows_from(decisions, ceilings, pref))
    yes_rows = [r for r in all_rows if r["Site-Specific?"] == "Yes"]

    if want_csv:
        for fn, rows in [("ICDO32_All_Screened_Terms.csv", all_rows),
                         ("ICDO32_SiteSpecific_Terms.csv", yes_rows)]:
            with open(fn, "w", newline="", encoding="utf-8") as f:
                w = csv.DictWriter(f, fieldnames=COLS); w.writeheader(); w.writerows(rows)

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    TEAL = "20808D"
    hdr_fill = PatternFill("solid", fgColor=TEAL)
    hdr_font = Font(name="Calibri", bold=True, color="FFFFFF")
    base_font = Font(name="Calibri")
    center = Alignment(horizontal="center", vertical="top", wrap_text=True)
    left = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # README
    ws = wb.active; ws.title = "README"
    readme = [
        "ICD-O-3.2 Site-Specific Term Detection - Full Catalog",
        "",
        "Method (locked):",
        "1. Ceiling: per code, allowed sites = SMVL rows with chosen status col == Valid (1). Unlikely/Impossible excluded.",
        "2. Terms verbatim from master ICDO32 list (code+term). Level & ICD-O suggested sites NOT used.",
        "3. Candidate universe = every term on a MULTI-site code. Single-site codes are auto 'No'.",
        "4. Literature confirmation: WHO Blue Books 5th ed, NCCN, UpToDate, CAP, PubMed. Every 'Yes' carries a source URL.",
        "5. Decision (any proper subset counts): 'Yes' when literature sites are a STRICT subset of the code's SMVL ceiling.",
        "",
        "Preferred Term column: master PrefTerm flag; /2 without preferred term falls back to /3 preferred marked '(for /3)'; otherwise blank.",
    ]
    for i, line in enumerate(readme):
        c = ws.cell(row=2+i, column=2, value=line)
        c.font = Font(name="Calibri", bold=(i in (0,2)), size=(14 if i==0 else 11))
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 120

    def build_sheet(title, subtitle, rows):
        ws = wb.create_sheet(title)
        ws.column_dimensions["A"].width = 3
        ws.cell(row=2, column=2, value=title).font = Font(name="Calibri", bold=True, size=14)
        ws.cell(row=3, column=2, value=subtitle).font = Font(name="Calibri", italic=True, size=10)
        for j, col in enumerate(COLS):
            c = ws.cell(row=5, column=2+j, value=col)
            c.fill = hdr_fill; c.font = hdr_font; c.alignment = center
        for i, row in enumerate(rows):
            for j, col in enumerate(COLS):
                c = ws.cell(row=6+i, column=2+j, value=row[col])
                c.font = base_font
                c.alignment = center if col in ("ICDO32 Code","Site-Specific?","Subset Site Count","Code's SMVL-Valid Site Count") else left
        widths = [12,34,30,12,26,30,10,14,34,40,50,40,34]
        from openpyxl.utils import get_column_letter
        for j, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(2+j)].width = w
        ws.freeze_panes = "B6"
        ws.auto_filter.ref = f"B5:{get_column_letter(1+len(COLS))}{5+len(rows)}"

    build_sheet("Site-Specific Terms", "Terms used at only a proper subset of their code's SMVL-valid sites.", yes_rows)
    build_sheet("All Screened Terms", "All screened terms with Yes/No/Uncertain decision and rationale.", all_rows)

    wb.save(out)
    print(f"wrote {out}: all={len(all_rows)} yes={len(yes_rows)}")

if __name__ == "__main__":
    main()
